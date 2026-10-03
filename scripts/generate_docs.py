"""Genera DOCUMENTACION_ARQUITECTURA_COMPLETA.docx a partir de DOCUMENTACION_ARQUITECTURA.md.
Uso: python scripts/generate_docs.py  (requiere: pip install python-docx)"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "DOCUMENTACION_ARQUITECTURA.md"
OUT = ROOT / "DOCUMENTACION_ARQUITECTURA_COMPLETA.docx"


def add_inline(par, text):
    # **negritas** y `código`
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            par.add_run(part)


def main():
    doc = Document()
    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            run = doc.add_paragraph().add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"-+", c) for c in cells):
                    rows.append(cells)
                i += 1
            i -= 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, cells in enumerate(rows):
                for c, text in enumerate(cells):
                    par = table.cell(r, c).paragraphs[0]
                    add_inline(par, text)
                    if r == 0:
                        for run in par.runs:
                            run.bold = True
        elif m := re.match(r"(#{1,4}) (.*)", line):
            doc.add_heading(m.group(2), level=len(m.group(1)) - 1 if len(m.group(1)) > 1 else 0)
        elif line.startswith("> "):
            add_inline(doc.add_paragraph(style="Intense Quote"), line[2:])
        elif m := re.match(r"(\s*)- (.*)", line):
            add_inline(doc.add_paragraph(style="List Bullet"), m.group(2))
        elif m := re.match(r"\d+\. (.*)", line):
            add_inline(doc.add_paragraph(style="List Number"), m.group(1))
        elif line.strip():
            add_inline(doc.add_paragraph(), line)
        i += 1
    doc.save(OUT)
    print(f"OK -> {OUT.name}")


if __name__ == "__main__":
    main()
